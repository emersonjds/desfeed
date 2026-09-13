"""Monta o roteiro de fala em DOCX a partir de `roteiro-pitch.md`.

Os papéis de parágrafo seguem o modelo de referência do haCARthon: contexto do evento,
título, abertura com o tempo-alvo, e por slide um trio — cabeçalho com faixa de tempo,
a fala, e a instrução de palco que termina na deixa para o próximo slide.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

AQUI = Path(__file__).parent

TINTA = RGBColor(0x0F, 0x13, 0x1D)
VERDE = RGBColor(0x05, 0x96, 0x69)
APOIO = RGBColor(0x4A, 0x54, 0x68)
FONTE = 'Helvetica Neue'

CONTEXTO = 'HACKTUDO 2026 · 13ª edição — PITCH (roteiro do speaker para ensaio)'
RODAPE = (
    'Roteiro de pitch — Memfeed · HACKTUDO 2026 · setembro de 2026. '
    'As fontes de cada número estão no rodapé da lâmina correspondente do deck.'
)


def criar_estilo(doc: Document, nome: str, tamanho: float, *, cor=TINTA,
                 negrito=False, italico=False, espaco_antes=0, espaco_depois=6):
    estilo = doc.styles.add_style(nome, 1)
    estilo.base_style = doc.styles['Normal']
    fonte = estilo.font
    fonte.name, fonte.size, fonte.bold, fonte.italic = FONTE, Pt(tamanho), negrito, italico
    fonte.color.rgb = cor
    estilo.paragraph_format.space_before = Pt(espaco_antes)
    estilo.paragraph_format.space_after = Pt(espaco_depois)
    return estilo


def escrever_com_negrito(paragrafo, texto: str, prefixo: str = ''):
    """`**assim**` no markdown vira negrito de verdade no Word."""
    if prefixo:
        run = paragrafo.add_run(prefixo)
        run.bold = True
        run.font.color.rgb = VERDE

    for pedaco in re.split(r'(\*\*[^*]+\*\*)', texto):
        if not pedaco:
            continue
        run = paragrafo.add_run(pedaco.strip('*'))
        run.bold = pedaco.startswith('**')


def montar() -> Path:
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = FONTE, Pt(11)

    criar_estilo(doc, 'MetaEvento', 8.5, cor=APOIO, negrito=True, espaco_depois=2)
    criar_estilo(doc, 'Abertura', 11, cor=APOIO, espaco_depois=8)
    criar_estilo(doc, 'CabecalhoSlide', 13.5, cor=VERDE, negrito=True,
                 espaco_antes=16, espaco_depois=4)
    criar_estilo(doc, 'Fala', 12.5, espaco_depois=5)
    criar_estilo(doc, 'Palco', 10, cor=APOIO, italico=True, espaco_depois=4)
    criar_estilo(doc, 'RodapeDoc', 8.5, cor=APOIO, espaco_antes=18)

    roteiro = (AQUI / 'roteiro-pitch.md').read_text()

    doc.add_paragraph(CONTEXTO, style='MetaEvento')
    titulo = doc.add_paragraph('Memfeed — Roteiro de Pitch', style='Title')
    titulo.runs[0].font.color.rgb = TINTA

    blocos = re.split(r'^## ', roteiro, flags=re.M)
    abertura = [l.strip() for l in blocos[0].split('\n')
                if l.strip() and not l.startswith('#') and not l.startswith('---')]
    for linha in abertura:
        escrever_com_negrito(doc.add_paragraph(style='Abertura'), linha.lstrip('- '))

    perguntas: list[tuple[str, str]] = []

    for bloco in blocos[1:]:
        linhas = bloco.split('\n')
        cabecalho = linhas[0].strip()

        if re.search(r'pergunta', cabecalho, re.I):
            perguntas = extrair_perguntas(linhas[1:])
            continue

        escrever_com_negrito(doc.add_paragraph(style='CabecalhoSlide'), cabecalho)

        for linha in linhas[1:]:
            texto = linha.strip()
            if not texto or texto.startswith('---'):
                continue
            if texto.startswith('> '):
                escrever_com_negrito(doc.add_paragraph(style='Fala'), texto[2:])
            elif texto != '>':
                texto = texto.replace('[TROCA]', '→ AVANÇA O SLIDE')
                escrever_com_negrito(doc.add_paragraph(style='Palco'), texto.lstrip('- '))

    if perguntas:
        escrever_com_negrito(doc.add_paragraph(style='CabecalhoSlide'),
                             'Perguntas prováveis dos jurados — respostas de bolso')
        tabela = doc.add_table(rows=1, cols=2)
        tabela.style = 'Light Grid Accent 1'
        for celula, titulo_col in zip(tabela.rows[0].cells, ('Pergunta', 'Resposta de bolso')):
            celula.text = ''
            run = celula.paragraphs[0].add_run(titulo_col)
            run.bold = True

        for pergunta, resposta in perguntas:
            linha = tabela.add_row().cells
            escrever_com_negrito(linha[0].paragraphs[0], pergunta)
            escrever_com_negrito(linha[1].paragraphs[0], resposta)

    rodape = doc.add_paragraph(RODAPE, style='RodapeDoc')
    rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER

    destino = AQUI / 'Memfeed-Roteiro-Speaker.docx'
    doc.save(destino)
    return destino


def extrair_perguntas(linhas: list[str]) -> list[tuple[str, str]]:
    """No markdown a pergunta é um `###` e a resposta é o parágrafo seguinte."""
    pares: list[tuple[str, str]] = []
    pergunta: str | None = None
    resposta: list[str] = []

    for linha in linhas:
        texto = linha.strip()
        if texto.startswith('### ') or re.match(r'^\*\*\d+\.', texto):
            if pergunta:
                pares.append((pergunta, ' '.join(resposta).strip()))
            pergunta = re.sub(r'^###\s*|^\*\*|\*\*$', '', texto).strip()
            resposta = []
        elif texto and not texto.startswith('---') and pergunta:
            resposta.append(texto.lstrip('> ').strip())

    if pergunta:
        pares.append((pergunta, ' '.join(resposta).strip()))
    return pares


if __name__ == '__main__':
    print(f'gerado: {montar()}')
