"""Monta o PPTX a partir dos slides renderizados do deck.html e das falas do roteiro.

O deck é desenhado em HTML/CSS. Redesenhar aquele layout com formas nativas do PowerPoint
custaria reimplementar o motor de layout e sairia pior, então cada slide entra como imagem
de página inteira e a fala correspondente vai para as notas do apresentador.
"""

import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches

AQUI = Path(__file__).parent

# 13,333 × 7,5 polegadas é o 16:9 padrão do PowerPoint, em EMU.
LARGURA = Emu(12192000)
ALTURA = Emu(6858000)

SLIDES_APRESENTADOS = 10

ANEXO_QA = 11


def _limpar(bloco: str) -> str:
    linhas = []
    for linha in bloco.split('\n'):
        limpa = linha.strip()
        if limpa.startswith('> '):
            limpa = limpa[2:]
        elif limpa == '>':
            limpa = ''
        limpa = re.sub(r'\*\*([^*]+)\*\*', r'\1', limpa)
        limpa = limpa.replace('[TROCA]', '→ avança o slide')
        linhas.append(limpa)
    return '\n'.join(linhas).strip()


def falas_por_slide(roteiro: str) -> dict[int, str]:
    """Extrai, de cada bloco `## Slide N`, a fala e a instrução de palco.

    O corte é em qualquer `##`: as seções finais do roteiro (versão curta, perguntas) não são
    slides e não podem grudar na nota do último.
    """
    blocos = re.split(r'^## ', roteiro, flags=re.M)[1:]
    falas: dict[int, str] = {}

    for bloco in blocos:
        if re.search(r'pergunta', bloco.split('\n')[0], re.I):
            falas[ANEXO_QA] = _limpar(bloco)
            continue

        numero = re.match(r'Slide (\d+)', bloco)
        if not numero:
            continue

        falas[int(numero.group(1))] = _limpar('\n'.join(bloco.split('\n')[1:]))

    return falas


def montar() -> Path:
    imagens = sorted(AQUI.glob('slides-png/hi-*.png'))
    if not imagens:
        raise SystemExit('renderize os slides antes: pdftoppm -png -r 200 deck.pdf slides-png/hi')

    falas = falas_por_slide((AQUI / 'roteiro-pitch.md').read_text())

    deck = Presentation()
    deck.slide_width = LARGURA
    deck.slide_height = ALTURA
    em_branco = deck.slide_layouts[6]

    for indice, imagem in enumerate(imagens, start=1):
        slide = deck.slides.add_slide(em_branco)
        slide.shapes.add_picture(str(imagem), Inches(0), Inches(0), width=LARGURA, height=ALTURA)

        nota = falas.get(indice, '')
        if indice > SLIDES_APRESENTADOS:
            nota = 'Anexo de Q&A — não conta como slide apresentado.\n\n' + nota
        slide.notes_slide.notes_text_frame.text = nota

    destino = AQUI / 'Memfeed-Pitch-Deck.pptx'
    deck.save(destino)
    return destino


if __name__ == '__main__':
    caminho = montar()
    print(f'gerado: {caminho}')
